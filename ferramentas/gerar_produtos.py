# Gera site/produtos_extra.json a partir da tabela ATACADO (aba COMPLETA): produtos que podem não estar
# na base do gerador, com descrição (PT), unidade, medidas, linha e status. Preço NÃO entra.
# Uso: python3 ferramentas/gerar_produtos.py "<ATACADO - TABELA COMPLETA E RESUMIDA.xlsx>"
import json, os, sys
import openpyxl

wb = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=True)
ws = wb["COMPLETA"] if "COMPLETA" in wb.sheetnames else wb.worksheets[0]
linhas = list(ws.iter_rows(values_only=True))
ih = next(i for i, r in enumerate(linhas) if r and "CODIGO" in [str(c or "").strip().upper() for c in r])
cab = [str(c or "").strip().upper() for c in linhas[ih]]
col = lambda nome: cab.index(nome)
c_st, c_cod, c_desc, c_med, c_um = col("STATUS"), col("CODIGO"), col("DESCRIÇÃO"), col("MEDIDAS"), col("1ª UNIDADE DE MEDIDA")
txt = lambda v: " ".join(str(v or "").split())
out, linha, vistos = [], "", set()
for r in linhas[ih + 1:]:
    if not r: continue
    cod, desc = txt(r[c_cod]), txt(r[c_desc])
    if not cod:                       # linha de título do grupo (ex.: "ACQUA KAP")
        if desc: linha = desc
        continue
    st = txt(r[c_st]).upper()
    if st == "FORA DE LINHA" or cod in vistos or not desc: continue
    vistos.add(cod)
    out.append({"cod": cod, "descricao": desc, "um": txt(r[c_um]).upper(), "medidas": txt(r[c_med]), "linha": linha, "status": st})
saida = os.path.join(os.path.dirname(__file__), "..", "site", "produtos_extra.json")
json.dump(out, open(saida, "w"), ensure_ascii=False, separators=(",", ":"))
print(f"{len(out)} produtos ({sum(p['status'] == 'LANÇAMENTO' for p in out)} lançamentos); {os.path.getsize(saida) // 1024} KB")
